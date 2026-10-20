#!/usr/bin/env python3
"""Ghép điểm đã tính + kết quả Agent 2 -> Excel và index.html (kiểu "CV Radar").

Dùng:  python -X utf8 tao_bao_cao.py <thu_muc_ket_qua>
Điều kiện: đã có criteria.json, scores.json, ket_qua_tinh_diem.json (chạy tinh_diem.py trước)
và verification_*.json hoặc verification.json (kết quả Agent 2).
Mẫu giao diện lấy từ ../assets (mau_bao_cao.html + giao_dien.js) hoặc từ thư mục kết quả.
index.html tạo ra TỰ CHỨA: ảnh CV nhúng dạng data URI nên mở được trong trình duyệt của
Claude Code / Codex (các trình duyệt này mở file cục bộ như trang "data:", không tải được ảnh tương đối).
"""
import json, os, re, base64, glob, difflib, io, sys, unicodedata
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

D = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
HERE = os.path.dirname(os.path.abspath(__file__))
J = lambda n: json.load(open(os.path.join(D, n), encoding="utf8"))


def tpl_path(name):
    for p in (os.path.join(D, name), os.path.join(HERE, "..", "assets", name), os.path.join(HERE, name)):
        if os.path.exists(p):
            return p
    raise FileNotFoundError(name)


crit, sc, kq = J("criteria.json"), J("scores.json"), J("ket_qua_tinh_diem.json")
TC = {t["id"]: t for t in crit["tieu_chi"]}
RES = {r["key"]: r for r in kq["ket_qua"]}
BAD = ("Không tìm thấy", "Sai ngữ cảnh")
VI_TRI = crit.get("vi_tri") or sc.get("vi_tri", "")
NGAY = sc["ngay"]
NGAY_VN = "/".join(reversed(NGAY.split("-")))
SLUG = re.sub(r"[^A-Za-z0-9]+", "_", unicodedata.normalize("NFD", VI_TRI.replace("đ", "d").replace("Đ", "D")).encode("ascii", "ignore").decode()).strip("_") or "vi_tri"
W_K = sum(t["trong_so"] for t in crit["tieu_chi"] if t["nhom"] == "K")
W_S = sum(t["trong_so"] for t in crit["tieu_chi"] if t["nhom"] == "S")

# ---------- ghép kết quả Agent 2 ----------
ver, other = [], []
for part in sorted(glob.glob(os.path.join(D, "verification_[A-Z].json"))):
    p = json.load(open(part, encoding="utf8"))
    ver += p.get("kiem_tra", [])
    other += p.get("kiem_tra_khac", [])
if ver:
    json.dump(ver, open(os.path.join(D, "verification.json"), "w", encoding="utf8"), ensure_ascii=False, indent=1)
    json.dump(other, open(os.path.join(D, "verification_nhan_xet.json"), "w", encoding="utf8"), ensure_ascii=False, indent=1)
elif os.path.exists(os.path.join(D, "verification.json")):
    ver = J("verification.json")
    other = J("verification_nhan_xet.json") if os.path.exists(os.path.join(D, "verification_nhan_xet.json")) else []
else:
    print("CẢNH BÁO: chưa có kết quả Agent 2, mọi trích dẫn sẽ ghi 'Chưa kiểm tra'.")

norm = lambda s: re.sub(r"\s+", " ", (s or "").lower()).strip(" .")


def find(key, tc, text, loai=None):
    best, bs = None, 0.0
    for v in ver:
        if v.get("key") != key:
            continue
        if tc and v.get("tieu_chi") != tc:
            continue
        if loai and v.get("loai") != loai:
            continue
        if loai is None and v.get("loai") == "Khang_dinh_khong_co":
            continue
        s = 1.0 if loai == "Khang_dinh_khong_co" else difflib.SequenceMatcher(None, norm(v.get("trich_dan")), norm(text)).ratio()
        if s > bs:
            best, bs = v, s
    return best if best and bs >= 0.6 else None


# ---------- dữ liệu ứng viên ----------
cands = []
for c in sc["ung_vien"]:
    r = RES[c["key"]]
    cells = []
    for i, t in TC.items():
        x = c["diem"][i]
        qs = []
        for q in x["trich_dan"]:
            v = find(c["key"], i, q["text"])
            qs.append({"text": q["text"], "vi_tri": q["vi_tri"], "ver": v["trang_thai"] if v else "Chưa kiểm tra", "ver_note": v.get("ghi_chu", "") if v else ""})
        av = None
        if not x["trich_dan"] and x["muc"] == 0:
            v = find(c["key"], i, "", "Khang_dinh_khong_co")
            av = {"ver": v["trang_thai"], "note": v.get("ghi_chu", "")} if v else None
        flag = any(q["ver"] in BAD for q in qs) or bool(av and av["ver"] in BAD)
        cells.append({"id": i, "nhom": t["nhom"], "mo_ta": t["mo_ta"], "muc_do": t["muc_do"], "trong_so": t["trong_so"], "muc": x["muc"],
                      "trang_thai": x["trang_thai"], "nhan_xet": x["nhan_xet"], "quotes": qs, "absence": av, "can_ra_soat": flag,
                      "dong_gop": round(t["trong_so"] * x["muc"] / 3, 2)})
    sig = []
    for a in c["thai_do"]["dau_hieu"]:
        v = find(c["key"], None, a["text"], "Thai_do")
        sig.append({"nhom": a["nhom"], "text": a["text"], "vi_tri": a["vi_tri"], "ver": v["trang_thai"] if v else "Chưa kiểm tra", "ver_note": v.get("ghi_chu", "") if v else ""})
    on = r["on_dinh"]
    on_txt = ("Có nhãn cần hỏi lại: " + "; ".join(f"{n} ({m} tháng)" for n, m in on["vi_tri_ngan"])) if on["nhan"] else \
        "Không có nhãn theo ngưỡng đã duyệt. " + ", ".join(f"{p['vai_tro']}: {p['so_thang']} tháng" + ("" if p["chinh_xac"] == "thang" else " (xấp xỉ, CV chỉ ghi năm)") for p in on["chi_tiet"])
    nflag = sum(1 for z in cells if z["can_ra_soat"]) + sum(1 for s in sig if s["ver"] in BAD)
    tt = c.get("tom_tat") or {}
    tin_hieu = tt.get("tin_hieu") or [m[:42] for m in c.get("diem_manh", [])[:4]]
    cands.append({"key": c["key"], "ten": c["ten"], "chuc_danh": c.get("chuc_danh_trong_cv", ""), "hang": r["hang"], "diem": r["diem_tong"], "K": r["diem_K"], "S": r["diem_S"],
                  "bao_phu": r["do_bao_phu"], "muc_phu_hop": r["muc_phu_hop"], "thieu_bb": r["tieu_chi_bat_buoc_thieu"], "cells": cells,
                  "thai_do": {"dau_hieu": sig, "star": c["thai_do"]["cau_hoi_star"], "on_dinh": on_txt},
                  "lam_ro": c["cau_hoi_lam_ro"], "manh": c.get("diem_manh", []), "luu_y": c.get("luu_y", []), "file": c["file_luu"],
                  "xem_truoc": c.get("xem_truoc") or c["file_luu"], "file_goc": c["file_goc"], "n_flag": nflag,
                  "tin_hieu": tin_hieu, "can_xac_minh": tt.get("can_xac_minh") or " · ".join(c.get("luu_y", [])[:2]),
                  "tong_ket": tt.get("tong_ket") or "; ".join(c.get("diem_manh", [])), "ly_do_hr": tt.get("ly_do_hr") or "; ".join(c.get("diem_manh", [])[:2])})
cands.sort(key=lambda z: (z["hang"], z["key"]))
allq = [q for z in cands for cell in z["cells"] for q in cell["quotes"]] + [s for z in cands for s in z["thai_do"]["dau_hieu"]]
cnt = {k: sum(1 for q in allq if q["ver"] == k) for k in ("Khớp", "Không tìm thấy", "Sai ngữ cảnh", "Chưa kiểm tra")}
n_abs = {k: sum(1 for z in cands for cell in z["cells"] if cell["absence"] and cell["absence"]["ver"] == k) for k in ("Khớp", "Không tìm thấy", "Sai ngữ cảnh")}
g_up = [z for z in cands if z["diem"] >= 70 and not z["thieu_bb"]]
g_dk = [z for z in cands if z not in g_up and z["diem"] >= 40]
g_hc = [z for z in cands if z not in g_up and z not in g_dk]
nhom_hr = [
    ("Ưu tiên mời phỏng vấn", g_up, "Điểm từ 70 trở lên và đủ mọi tiêu chí bắt buộc. Vẫn cần xác minh các điểm ở cột 'Kiểm tra trước quyết định'."),
    ("Xem xét có điều kiện, cần làm rõ trước", g_dk, "Có bằng chứng một phần nhưng thiếu tiêu chí bắt buộc hoặc điểm chưa tới 70; làm rõ trước khi quyết định."),
    ("Bằng chứng còn hạn chế", g_hc, "Thiếu hầu hết tiêu chí bắt buộc hoặc thuộc lĩnh vực khác. Chỉ là gợi ý rà soát, HR quyết định."),
]
json.dump({"nhom_hr": [(a, [z["key"] for z in b], c) for a, b, c in nhom_hr], "dem": cnt, "vang_mat": n_abs}, open(os.path.join(D, "tong_hop.json"), "w", encoding="utf8"), ensure_ascii=False, indent=1)

# ======================= EXCEL =======================
NAVY = "17375E"
HDR = PatternFill("solid", fgColor=NAVY)
HF = Font(bold=True, color="FFFFFF", name="Arial", size=10)
GRN = PatternFill("solid", fgColor="E6F4EA"); RED = PatternFill("solid", fgColor="FDE7E7"); BLU = PatternFill("solid", fgColor="DDEEFB")
YEL = PatternFill("solid", fgColor="FFF4CC"); LBL = PatternFill("solid", fgColor="E7EEF8")
F = lambda **k: Font(name="Arial", size=10, **k)
thin = Side(style="thin", color="D9D9D9"); BRD = Border(bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")


def ts(c):
    return "Đáp ứng" if c["muc"] >= 2 else "Đáp ứng một phần" if c["muc"] == 1 else ("Không đáp ứng" if c["trang_thai"] == "Không đáp ứng" else "Chưa đủ bằng chứng")


def lab(v):
    return {"Khớp": "Khớp", "Không tìm thấy": "KHÔNG TÌM THẤY", "Sai ngữ cảnh": "SAI NGỮ CẢNH"}.get(v, "Chưa kiểm tra")


def style_level(cell, txt):
    cell.fill = GRN if ("rất mạnh" in txt or "tốt" in txt) else BLU if "một phần" in txt else RED
    cell.font = F(bold=True, color=("1E7B34" if cell.fill == GRN else "1F5F99" if cell.fill == BLU else "B42318"))


def hdr(ws, row, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.fill, c.font, c.alignment = HDR, HF, Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def widths(ws, ws_):
    for i, w in enumerate(ws_, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def body(ws, r0, rows):
    for k, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r0 + k, column=j, value=v)
            c.alignment, c.font, c.border = WR, F(), BRD


wb = Workbook()
ws = wb.active
ws.title = "Tổng quan"
ws["A1"] = f"Sàng lọc CV · {VI_TRI}"
ws["A1"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
meta = [("Số CV đã đọc", len(cands)), ("Nguồn", sc.get("nguon", "")), ("Phạm vi", sc.get("pham_vi", f"{len(cands)} file CV")),
        ("Tiêu chí ASK", f"K {W_K}% · S {W_S}% · A không chấm điểm (chỉ hành vi cần xác minh)"),
        ("Lưu ý", "Điểm là mức bằng chứng K+S trong CV. Không phải quyết định tuyển/loại. " + " ".join(crit.get("gia_dinh", [])))]
for i, (a, b) in enumerate(meta, 2):
    ws.cell(row=i, column=1, value=a).font = F(bold=True, color=NAVY)
    c = ws.cell(row=i, column=2, value=b)
    c.font = F(); c.alignment = Alignment(wrap_text=False, vertical="top", horizontal="left")
r0 = 8
hdr(ws, r0, ["Nhóm HR xem kỹ", "Ứng viên", "Lý do", "Kiểm tra trước quyết định"])
rows = []
for gi, (gname, zs, why) in enumerate(nhom_hr[:2]):
    for z in zs:
        rows.append([gname if gi == 0 else "Có điều kiện", z["ten"], z["ly_do_hr"], z["can_xac_minh"]])
body(ws, r0 + 1, rows)
r = r0 + 1 + len(rows) + 1
for j, v in enumerate(["Đánh giá từ CV", "Đáp ứng", "Đáp ứng một phần", "Chưa đủ bằng chứng"], 1):
    c = ws.cell(row=r, column=j, value=v); c.fill, c.font = LBL, F(bold=True, color=NAVY)
body(ws, r + 1, [["Thang điểm 0–3", "Bằng chứng trực tiếp (mức 2–3)", "Bằng chứng gián tiếp/chưa đủ phạm vi (mức 1)", "CV không nêu (mức 0). Không suy luận ứng viên không có năng lực."]])
r += 3
for j, v in enumerate(["Kiểm tra độc lập (Agent 2)", "Khớp", "Sai ngữ cảnh", "Không tìm thấy"], 1):
    c = ws.cell(row=r, column=j, value=v); c.fill, c.font = LBL, F(bold=True, color=NAVY)
body(ws, r + 1, [["Số câu trích đã kiểm tra", cnt["Khớp"], cnt["Sai ngữ cảnh"], cnt["Không tìm thấy"]],
                 ["Khẳng định 'CV không nêu'", n_abs["Khớp"], n_abs["Sai ngữ cảnh"], n_abs["Không tìm thấy"]]])
ws.cell(row=r + 2, column=3).fill = RED if cnt["Sai ngữ cảnh"] else GRN
r += 4
ws.cell(row=r, column=1, value="Các điểm Agent 2 đề nghị xem lại (chưa sửa điểm, HR quyết định)").font = F(bold=True, color=NAVY)
nl = [o for o in other if "không phát hiện" not in (o.get("van_de", "") + o.get("ghi_chu", "")).lower() and o.get("key") != "toan_cuc"]
body(ws, r + 1, [[o.get("ung_vien", ""), o.get("van_de", ""), o.get("ghi_chu", ""), ""] for o in nl])
widths(ws, [30, 38, 70, 62])

ws2 = wb.create_sheet("Ứng viên")
hdr(ws2, 1, ["Hạng", "Ứng viên", "Điểm bằng chứng", "Độ bao phủ", "Điểm S", "Điểm K", "Mức bằng chứng", "Tiêu chí bắt buộc còn thiếu", "Kiểm tra độc lập", "CV gốc"])
for i, z in enumerate(cands, 2):
    vals = [z["hang"], z["ten"], z["diem"], z["bao_phu"] / 100, z["S"], z["K"], z["muc_phu_hop"], ", ".join(z["thieu_bb"]) or "—",
            f'{z["n_flag"]} mục cần rà soát' if z["n_flag"] else "Khớp", z["file_goc"]]
    for j, v in enumerate(vals, 1):
        c = ws2.cell(row=i, column=j, value=v)
        c.font, c.alignment, c.border = F(bold=(j == 2)), Alignment(vertical="center", wrap_text=True), BRD
    ws2.cell(row=i, column=3).number_format = "0.0"; ws2.cell(row=i, column=4).number_format = "0%"
    for j in (5, 6):
        ws2.cell(row=i, column=j).number_format = "0.0"
    style_level(ws2.cell(row=i, column=7), z["muc_phu_hop"])
    if z["thieu_bb"]:
        ws2.cell(row=i, column=8).fill = YEL
    if z["n_flag"]:
        ws2.cell(row=i, column=9).fill = RED; ws2.cell(row=i, column=9).font = F(bold=True, color="B42318")
    lk = ws2.cell(row=i, column=10); lk.hyperlink = z["file"]; lk.font = F(color="0563C1", underline="single")
    ws2.row_dimensions[i].height = 22
ws2.freeze_panes = "C2"
widths(ws2, [7, 24, 16, 12, 10, 10, 30, 26, 24, 28])

ws3 = wb.create_sheet("Ma trận đối chiếu")
ws3["A1"] = "Ma trận đối chiếu tiêu chí ASK"; ws3["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
ws3["A2"] = "Mỗi hàng là một ứng viên × một tiêu chí. Điểm quy đổi = trọng số × mức / 3 (tính bằng tinh_diem.py). Ô đỏ = Agent 2 gắn Không tìm thấy hoặc Sai ngữ cảnh."
ws3["A2"].font = F(italic=True, color="666666")
hdr(ws3, 4, ["Ứng viên", "ID", "ASK", "Tiêu chí", "Mức độ", "Trọng số (%)", "Điểm (0–3)", "Điểm quy đổi", "Trạng thái", "Bằng chứng từ CV (nguyên văn)", "Nguồn CV (trang, mục)", "Kiểm tra độc lập", "Câu hỏi cần xác minh"])
rows = []
for z in cands:
    for c in z["cells"]:
        qt = "\n".join(f'“{q["text"]}”' for q in c["quotes"]) or c["nhan_xet"]
        vt_ = "\n".join(q["vi_tri"] for q in c["quotes"]) or "—"
        if c["quotes"]:
            vv = "\n".join(lab(q["ver"]) + (f' ({q["ver_note"]})' if q["ver"] in BAD and q["ver_note"] else "") for q in c["quotes"])
        elif c["absence"]:
            vv = (lab(c["absence"]["ver"]) + " (đã xác nhận CV không nêu)") if c["absence"]["ver"] == "Khớp" else (lab(c["absence"]["ver"]) + " " + c["absence"]["note"])
        else:
            vv = "—"
        rows.append([z["ten"], c["id"], c["nhom"], c["mo_ta"], c["muc_do"], c["trong_so"], c["muc"], c["dong_gop"], ts(c), qt, vt_, vv, "; ".join(z["lam_ro"][:1])])
body(ws3, 5, rows)
for k, row in enumerate(rows, 5):
    st = row[8]; cell = ws3.cell(row=k, column=9)
    cell.fill = GRN if st == "Đáp ứng" else BLU if st == "Đáp ứng một phần" else RED
    cell.font = F(color=("1E7B34" if st == "Đáp ứng" else "1F5F99" if st == "Đáp ứng một phần" else "B42318"))
    if "KHÔNG TÌM THẤY" in row[11] or "SAI NGỮ CẢNH" in row[11]:
        for j in (10, 11, 12):
            ws3.cell(row=k, column=j).fill = RED
    ws3.cell(row=k, column=8).number_format = "0.00"
ws3.freeze_panes = "B5"
widths(ws3, [22, 6, 6, 44, 10, 11, 10, 12, 18, 70, 36, 30, 44])

ws4 = wb.create_sheet("Thái độ")
ws4["A1"] = "Thái độ (Attitude): không chấm điểm, không trọng số, không dùng xếp hạng"; ws4["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
hdr(ws4, 3, ["Ứng viên", "Dấu hiệu hành vi trong CV (trích dẫn, vị trí, kiểm tra độc lập)", "Dấu hiệu ổn định công việc (dữ kiện để hỏi lại)", "Câu hỏi phỏng vấn hành vi (STAR)"])
rows = []
for z in cands:
    sg = "\n".join(f'[{s["nhom"]}] “{s["text"]}” ({s["vi_tri"]}) → {lab(s["ver"])}' for s in z["thai_do"]["dau_hieu"]) or "Không tìm thấy dấu hiệu"
    rows.append([z["ten"], sg, z["thai_do"]["on_dinh"], "\n".join(f"{i}. {q}" for i, q in enumerate(z["thai_do"]["star"], 1))])
body(ws4, 4, rows)
for k, row in enumerate(rows, 4):
    if "KHÔNG TÌM THẤY" in row[1] or "SAI NGỮ CẢNH" in row[1]:
        ws4.cell(row=k, column=2).fill = RED
widths(ws4, [24, 90, 55, 80]); ws4.freeze_panes = "B4"

ws5 = wb.create_sheet("Hướng dẫn ASK")
ws5.column_dimensions["A"].width = 120
guide = ["ASK là gì?", "A (Attitude, Thái độ): hành vi cần xác minh khi phỏng vấn. KHÔNG chấm điểm, không trọng số, không đưa vào xếp hạng.",
         f"S (Skills, Kỹ năng): việc đã làm, công cụ, số năm, kết quả. Trọng số {W_S}/100.", f"K (Knowledge, Kiến thức): học vấn, kiến thức chuyên môn, chứng chỉ. Trọng số {W_K}/100.", "",
         "Thang 0–3: 3 mạnh (trực tiếp, có số liệu hoặc lặp ở nhiều vai trò) · 2 đủ · 1 một phần (gián tiếp) · 0 không tìm thấy.",
         f"Công thức: Điểm bằng chứng = Σ(trọng số × mức / 3). Điểm K, S quy về thang 100 của từng nhóm. Độ bao phủ = số tiêu chí có mức > 0 / {len(TC)}.",
         "Mức diễn giải: 85–100 rất mạnh · 70–84,9 tốt · 50–69,9 một phần · dưới 50 hạn chế. Là nhãn hỗ trợ rà soát, không phải quyết định tuyển/loại.",
         "Điểm 0 chỉ có nghĩa CV không nêu, không kết luận người đó thiếu năng lực. Tiêu chí bắt buộc còn thiếu được tô vàng.",
         "Giới hạn: đánh giá chỉ dựa trên nội dung CV; không xác minh với nhà tuyển dụng cũ; mốc chỉ ghi năm được tính xấp xỉ.", "Giả định: " + " ".join(crit.get("gia_dinh", []))]
for g in guide:
    ws5.append([g]); ws5[ws5.max_row][0].alignment = WR; ws5[ws5.max_row][0].font = F()
ws5["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
xlsx = os.path.join(D, f"Danh_gia_CV_{SLUG}.xlsx")
wb.save(xlsx)

# ======================= HTML (tự chứa) =======================
def jpeg_uri(path, maxw=1100):
    full = os.path.join(D, path)
    try:
        im = Image.open(full).convert("RGB")
    except Exception:  # DOCX/định dạng không xem trước được
        from PIL import ImageDraw
        im = Image.new("RGB", (900, 360), "#f2f6f6")
        ImageDraw.Draw(im).text((30, 160), "Khong co anh xem truoc. Mo file CV goc trong assets/cv de xem.", fill="#365c52")
    if im.width > maxw:
        im = im.resize((maxw, int(im.height * maxw / im.width)), Image.LANCZOS)
    b = io.BytesIO(); im.save(b, "JPEG", quality=80, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(b.getvalue()).decode()


PRI = {1: "Ưu tiên 1", 2: "Ưu tiên 2", 3: "Ưu tiên 3"}
top3 = [z["key"] for z in cands[:3]]
out = []
for z in cands:
    crit_list = [{"id": c["id"], "group": c["nhom"], "text": c["mo_ta"], "priority": c["muc_do"], "score": c["muc"], "evidence": c["nhan_xet"],
                  "quotes": c["quotes"], "absence": c["absence"], "flag": c["can_ra_soat"]} for c in z["cells"]]
    out.append({"key": z["key"], "name": z["ten"], "role": z["chuc_danh"], "total": z["diem"], "coverage": z["bao_phu"] / 100, "S": z["S"], "K": z["K"], "level": z["muc_phu_hop"],
                "priority": PRI[top3.index(z["key"]) + 1] if z["key"] in top3 else None, "signals": z["tin_hieu"], "verify": z["can_xac_minh"], "summary": z["tong_ket"],
                "checks": z["luu_y"][:4], "missing": z["thieu_bb"], "criteria": crit_list, "questions": z["lam_ro"],
                "att": {"signals": z["thai_do"]["dau_hieu"], "star": z["thai_do"]["star"], "stability": z["thai_do"]["on_dinh"]}, "image": jpeg_uri(z["xem_truoc"])})
tpl = open(tpl_path("mau_bao_cao.html"), encoding="utf8").read()
js = open(tpl_path("giao_dien.js"), encoding="utf8").read()
mo_ta = crit.get("mo_ta_nhom", {})
tong = sum(cnt.values())
esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;")
html = (tpl.replace("__DATA__", json.dumps(out, ensure_ascii=False).replace("</", "<\\/")).replace("__EMBED__", "true")
        .replace("__JS__", "const DATAVER=" + json.dumps({"khop": cnt["Khớp"], "tong": tong, "wK": W_K, "wS": W_S}) + ";\n" + js)
        .replace("__VI_TRI__", esc(VI_TRI)).replace("__NGAY_VN__", NGAY_VN).replace("__N__", str(len(cands)))
        .replace("__NGUON__", esc(sc.get("nguon", "Thư mục CV do người dùng cung cấp")))
        .replace("__ASK_A__", esc(mo_ta.get("A", "Hành vi cần xác minh khi phỏng vấn. Chỉ ghi dấu hiệu trong CV và câu hỏi STAR, không có điểm, không dùng xếp hạng.")))
        .replace("__ASK_S__", esc(mo_ta.get("S", "Kỹ năng thực hành: công cụ, số năm kinh nghiệm, kết quả đo được.")))
        .replace("__ASK_K__", esc(mo_ta.get("K", "Kiến thức nền: học vấn, chứng chỉ, hiểu biết lĩnh vực.")))
        .replace("__W_K__", str(W_K)).replace("__W_S__", str(W_S)).replace("__GIA_DINH__", esc(" ".join(crit.get("gia_dinh", [])))))
open(os.path.join(D, "index.html"), "w", encoding="utf8").write(html)
print("Xong:", cnt, n_abs, "| index.html", round(len(html) / 1e6, 2), "MB |", os.path.basename(xlsx))
